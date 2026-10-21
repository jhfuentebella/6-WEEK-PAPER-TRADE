from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
from typing import Optional

import pandas as pd
from openpyxl import load_workbook


PROJECT_YEAR = 2026


def normalize_ticker(code):
    if code is None:
        return None
    s = str(code).strip()
    if not s or s in {"0", "None", "nan"}:
        return None

    # Course files prefix KRX numeric codes with A to preserve them as text.
    if re.fullmatch(r"A\d{6}", s):
        return s[1:] + ".KS"
    if re.fullmatch(r"\d{6}", s):
        return s + ".KS"

    return s


def _cell_text(v):
    return "" if v is None else str(v).strip()


def _find_row(ws, needles, start=1, end=None):
    if isinstance(needles, str):
        needles = [needles]
    needles = [n.lower() for n in needles]
    end = end or ws.max_row

    for r in range(start, end + 1):
        for c in range(1, ws.max_column + 1):
            t = _cell_text(ws.cell(r, c).value).lower()
            if any(n in t for n in needles):
                return r
    return None


def _find_exact_row(ws, labels, start=1, end=None):
    if isinstance(labels, str):
        labels = [labels]
    labels = {x.lower() for x in labels}
    end = end or ws.max_row

    for r in range(start, end + 1):
        for c in range(1, ws.max_column + 1):
            if _cell_text(ws.cell(r, c).value).lower() in labels:
                return r
    return None


def _sheet_date(ws, filename=""):
    # Valuation sheets place the as-of date around K1.
    for r in range(1, min(ws.max_row, 6) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, datetime):
                return pd.Timestamp(v).normalize()
            if isinstance(v, str):
                m = re.fullmatch(r"\d{4}-\d{2}-\d{2}", v.strip())
                if m:
                    return pd.Timestamp(v).normalize()

    # Try Trade Date row.
    for r in range(1, min(ws.max_row, 10) + 1):
        row_vals = [_cell_text(ws.cell(r, c).value).lower() for c in range(1, ws.max_column + 1)]
        if "trade date" in row_vals:
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if isinstance(v, datetime):
                    return pd.Timestamp(v).normalize()

    # Fall back to sheet name, e.g. "16-Oct".
    m = re.fullmatch(r"(\d{1,2})-([A-Za-z]{3})", ws.title.strip())
    if m:
        day = int(m.group(1))
        mon = m.group(2).title()
        year_match = re.search(r"(20\d{2})\d{4}", filename)
        year = int(year_match.group(1)) if year_match else PROJECT_YEAR
        return pd.Timestamp(f"{year}-{mon}-{day}")

    return None


def _row_values_by_code_columns(ws, code_row, other_rows):
    records = []
    for c in range(1, ws.max_column + 1):
        ticker = normalize_ticker(ws.cell(code_row, c).value)
        if ticker is None:
            continue

        # Exclude text labels accidentally interpreted as tickers.
        if ticker.lower() in {
            "product code", "product name", "product type (etf/futures/options)",
            "size (shares/contracts)", "size (shares, contracts)"
        }:
            continue

        rec = {"ticker": ticker, "_col": c}
        for key, row in other_rows.items():
            rec[key] = ws.cell(row, c).value if row else None
        records.append(rec)
    return records


def _find_label_row_after(ws, labels, start, end=None):
    return _find_exact_row(ws, labels, start=start, end=end or ws.max_row)


def parse_positions_block(ws, block_row, sheet_date):
    end = min(ws.max_row, block_row + 20)

    side_row = _find_label_row_after(ws, ["Long or Short"], block_row + 1, end)
    name_row = _find_label_row_after(ws, ["Product Name"], block_row + 1, end)
    code_row = _find_label_row_after(ws, ["Product Code"], block_row + 1, end)
    size_row = _find_label_row_after(
        ws, ["Size (shares/contracts)", "Size (shares, contracts)"], block_row + 1, end
    )

    price_row = _find_row(
        ws,
        ["booking/opening price", "new price ("],
        start=block_row + 1,
        end=end,
    )
    fx_row = _find_row(
        ws,
        ["booking/opening fx", "new fx (exchange rate)"],
        start=block_row + 1,
        end=end,
    )
    mv_row = _find_exact_row(ws, ["Market Values of ETFs"], block_row + 1, end)

    if not code_row or not size_row:
        return []

    rows = _row_values_by_code_columns(
        ws,
        code_row,
        {
            "side": side_row,
            "name": name_row,
            "quantity": size_row,
            "recorded_price": price_row,
            "recorded_fx": fx_row,
            "recorded_market_value": mv_row,
        },
    )

    out = []
    for x in rows:
        try:
            qty = float(x["quantity"])
        except Exception:
            continue
        if qty == 0:
            continue

        side = _cell_text(x["side"]) or "Long"
        out.append(
            {
                "date": sheet_date,
                "ticker": x["ticker"],
                "name": _cell_text(x["name"]) or x["ticker"],
                "side": side.title(),
                "quantity": abs(qty),
                "recorded_price": pd.to_numeric(x["recorded_price"], errors="coerce"),
                "recorded_fx": pd.to_numeric(x["recorded_fx"], errors="coerce"),
                "recorded_market_value": pd.to_numeric(x["recorded_market_value"], errors="coerce"),
            }
        )
    return out


def parse_trades_block(ws, block_row, sheet_date):
    end = min(ws.max_row, block_row + 18)

    action_row = _find_label_row_after(ws, ["Buy or Sell"], block_row + 1, end)
    name_row = _find_label_row_after(ws, ["Product Name"], block_row + 1, end)
    code_row = _find_label_row_after(ws, ["Product Code"], block_row + 1, end)
    size_row = _find_label_row_after(
        ws, ["Size (shares/contracts)", "Size (shares, contracts)"], block_row + 1, end
    )
    price_row = _find_row(ws, ["new price ("], start=block_row + 1, end=end)
    fx_row = _find_row(ws, ["new fx (exchange rate)"], start=block_row + 1, end=end)

    if not code_row or not size_row or not action_row:
        return []

    rows = _row_values_by_code_columns(
        ws,
        code_row,
        {
            "action": action_row,
            "name": name_row,
            "quantity": size_row,
            "recorded_price": price_row,
            "recorded_fx": fx_row,
        },
    )

    out = []
    for x in rows:
        action = _cell_text(x["action"]).upper()
        if action not in {"BUY", "SELL"}:
            continue
        try:
            qty = float(x["quantity"])
        except Exception:
            continue
        if qty == 0:
            continue
        out.append(
            {
                "date": sheet_date,
                "ticker": x["ticker"],
                "name": _cell_text(x["name"]) or x["ticker"],
                "action": action,
                "quantity": abs(qty),
                "recorded_price": pd.to_numeric(x["recorded_price"], errors="coerce"),
                "recorded_fx": pd.to_numeric(x["recorded_fx"], errors="coerce"),
                "source": "Valuation",
            }
        )
    return out


def parse_order_block(ws, block_row, sheet_date):
    end = min(ws.max_row, block_row + 15)

    action_row = _find_label_row_after(ws, ["Buy or Sell"], block_row + 1, end)
    name_row = _find_label_row_after(ws, ["Product Name"], block_row + 1, end)
    code_row = _find_label_row_after(ws, ["Product Code"], block_row + 1, end)
    size_row = _find_label_row_after(
        ws, ["Size (shares/contracts)", "Size (shares, contracts)"], block_row + 1, end
    )

    if not code_row or not size_row or not action_row:
        return []

    rows = _row_values_by_code_columns(
        ws,
        code_row,
        {"action": action_row, "name": name_row, "quantity": size_row},
    )

    out = []
    for x in rows:
        action = _cell_text(x["action"]).upper()
        if action not in {"BUY", "SELL"}:
            continue
        try:
            qty = float(x["quantity"])
        except Exception:
            continue
        if qty == 0:
            continue
        out.append(
            {
                "date": sheet_date,
                "ticker": x["ticker"],
                "name": _cell_text(x["name"]) or x["ticker"],
                "action": action,
                "quantity": abs(qty),
                "source": "Order",
            }
        )
    return out


def _extract_cash(ws):
    row = _find_exact_row(ws, ["Nextweek Cash Balance"])
    if row:
        # In the course template, the value sits directly above the label in the same column.
        for c in range(1, ws.max_column + 1):
            if _cell_text(ws.cell(row, c).value).lower() == "nextweek cash balance":
                v = ws.cell(row - 1, c).value
                v = pd.to_numeric(v, errors="coerce")
                if pd.notna(v):
                    return float(v)

    row = _find_exact_row(ws, ["New Cash Balance"])
    if row:
        for c in range(1, ws.max_column + 1):
            if _cell_text(ws.cell(row, c).value).lower() == "new cash balance":
                v = ws.cell(row - 1, c).value
                v = pd.to_numeric(v, errors="coerce")
                if pd.notna(v):
                    return float(v)

    return None


def _extract_nav_return(ws, is_initial=False):
    row = _find_exact_row(ws, ["New Cash Balance + Unrealized P/L"])
    if row:
        nums = []
        for c in range(1, ws.max_column + 1):
            v = pd.to_numeric(ws.cell(row, c).value, errors="coerce")
            if pd.notna(v):
                nums.append(float(v))
        nav = next((x for x in nums if abs(x) > 1_000_000), None)
        ret = next((x for x in nums if -2 < x < 2), None)
        return nav, ret

    if is_initial:
        # Initial valuation is the KRW 10bn starting portfolio.
        return 10_000_000_000.0, 0.0

    return None, None


def parse_valuation_workbook(file_bytes: bytes, filename: str = "valuation.xlsx"):
    wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=False)

    snapshots = []
    all_trades = []

    for ws in wb.worksheets:
        sheet_date = _sheet_date(ws, filename)
        if sheet_date is None:
            continue

        block_row = _find_exact_row(ws, ["Outstanding (New) Positions"])
        is_initial = False
        if not block_row:
            block_row = _find_exact_row(ws, ["New Positions"])
            is_initial = block_row is not None

        if not block_row:
            continue

        positions = parse_positions_block(ws, block_row, sheet_date)
        if not positions:
            continue

        trade_block = _find_exact_row(ws, ["New Trades"])
        trades = parse_trades_block(ws, trade_block, sheet_date) if trade_block else []
        all_trades.extend(trades)

        cash = _extract_cash(ws)
        nav, ret = _extract_nav_return(ws, is_initial=is_initial)

        # If initial cash is just the tiny residual and no explicit NAV exists, NAV is still 10bn.
        snapshots.append(
            {
                "date": sheet_date,
                "sheet": ws.title,
                "positions": positions,
                "cash": cash,
                "nav": nav,
                "return_from_inception": ret,
            }
        )

    snapshots.sort(key=lambda x: x["date"])
    trades_df = pd.DataFrame(all_trades)
    if not trades_df.empty:
        trades_df = trades_df.sort_values(["date", "ticker"]).reset_index(drop=True)

    return {
        "filename": filename,
        "snapshots": snapshots,
        "trades": trades_df,
        "latest": snapshots[-1] if snapshots else None,
    }


def parse_order_workbook(file_bytes: bytes, filename: str = "orders.xlsx"):
    wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=False)

    all_orders = []
    sheet_dates = []

    for ws in wb.worksheets:
        sheet_date = _sheet_date(ws, filename)
        if sheet_date is None:
            continue

        block_row = _find_row(ws, ["moc(market-on-close) orders"])
        if not block_row:
            continue

        orders = parse_order_block(ws, block_row, sheet_date)
        if orders:
            all_orders.extend(orders)
            sheet_dates.append(sheet_date)

    orders_df = pd.DataFrame(all_orders)
    if not orders_df.empty:
        orders_df = orders_df.sort_values(["date", "ticker"]).reset_index(drop=True)

    return {
        "filename": filename,
        "orders": orders_df,
        "latest_order_date": max(sheet_dates) if sheet_dates else None,
    }


def first_acquisition_dates(valuation_result, order_result):
    dates = {}

    for snap in valuation_result.get("snapshots", []):
        for p in snap["positions"]:
            t = p["ticker"]
            dates[t] = min(dates.get(t, snap["date"]), snap["date"])

    orders = order_result.get("orders", pd.DataFrame())
    if not orders.empty:
        for _, r in orders[orders["action"] == "BUY"].iterrows():
            t = r["ticker"]
            d = pd.Timestamp(r["date"])
            dates[t] = min(dates.get(t, d), d)

    return dates


def latest_adjustments(valuation_result, order_result):
    adjustments = {}

    # Authoritative completed trades from valuation.
    trades = valuation_result.get("trades", pd.DataFrame())
    if not trades.empty:
        grouped = trades.groupby(["date", "ticker"], sort=True)
        for (d, t), g in grouped:
            buy_qty = g.loc[g["action"] == "BUY", "quantity"].sum()
            sell_qty = g.loc[g["action"] == "SELL", "quantity"].sum()
            if sell_qty > 0 and buy_qty == 0:
                text = f"{pd.Timestamp(d):%Y-%m-%d} • Sold {sell_qty:,.0f}"
            elif buy_qty > 0 and sell_qty == 0:
                text = f"{pd.Timestamp(d):%Y-%m-%d} • Bought +{buy_qty:,.0f}"
            else:
                text = f"{pd.Timestamp(d):%Y-%m-%d} • Rebalanced"
            adjustments[t] = (pd.Timestamp(d), text)

    # Submitted orders only matter if later than the latest valuation snapshot date.
    latest_val_date = (
        valuation_result["latest"]["date"]
        if valuation_result.get("latest")
        else pd.Timestamp.min
    )
    orders = order_result.get("orders", pd.DataFrame())
    if not orders.empty:
        later = orders[orders["date"] > latest_val_date]
        for (d, t), g in later.groupby(["date", "ticker"], sort=True):
            buy_qty = g.loc[g["action"] == "BUY", "quantity"].sum()
            sell_qty = g.loc[g["action"] == "SELL", "quantity"].sum()
            if sell_qty > 0 and buy_qty == 0:
                text = f"{pd.Timestamp(d):%Y-%m-%d} • Submitted sell {sell_qty:,.0f}"
            elif buy_qty > 0 and sell_qty == 0:
                text = f"{pd.Timestamp(d):%Y-%m-%d} • Submitted buy +{buy_qty:,.0f}"
            else:
                text = f"{pd.Timestamp(d):%Y-%m-%d} • Submitted rebalance"
            adjustments[t] = (pd.Timestamp(d), text)

    # Positions with no trade on the latest order date are "Retain".
    latest_order_date = order_result.get("latest_order_date")
    latest = valuation_result.get("latest")
    if latest_order_date is not None and latest is not None and latest_order_date >= latest["date"]:
        ordered_tickers = set()
        if not orders.empty:
            ordered_tickers = set(orders.loc[orders["date"] == latest_order_date, "ticker"])
        for p in latest["positions"]:
            t = p["ticker"]
            if t not in ordered_tickers:
                current = adjustments.get(t)
                if current is None or current[0] < latest_order_date:
                    adjustments[t] = (
                        latest_order_date,
                        f"{latest_order_date:%Y-%m-%d} • Retain",
                    )

    return {k: v[1] for k, v in adjustments.items()}


def current_position_state(valuation_result, order_result):
    """
    Start from the latest Monday valuation (authoritative).
    If a newer Friday order exists, apply its quantities as an estimate until the
    next Monday valuation is uploaded.
    """
    latest = valuation_result.get("latest")
    if latest is None:
        return {}, 0.0, "No valuation uploaded"

    state = {}
    for p in latest["positions"]:
        sign = -1 if p["side"].lower() == "short" else 1
        state[p["ticker"]] = {
            "ticker": p["ticker"],
            "name": p["name"],
            "signed_quantity": sign * float(p["quantity"]),
            "side": p["side"],
            "valuation_price": p["recorded_price"],
            "valuation_fx": p["recorded_fx"],
            "valuation_date": latest["date"],
        }

    cash = float(latest["cash"] or 0.0)
    mode = f"Authoritative valuation as of {latest['date']:%Y-%m-%d}"

    orders = order_result.get("orders", pd.DataFrame())
    if not orders.empty:
        pending = orders[orders["date"] > latest["date"]]
        if not pending.empty:
            # Holdings quantities can be updated immediately from submitted order quantities.
            # Cash is left at the valuation cash until online trade-price estimation is applied in app.py.
            for _, r in pending.sort_values("date").iterrows():
                t = r["ticker"]
                if t not in state:
                    state[t] = {
                        "ticker": t,
                        "name": r["name"],
                        "signed_quantity": 0.0,
                        "side": "Long",
                        "valuation_price": None,
                        "valuation_fx": None,
                        "valuation_date": latest["date"],
                    }

                q = state[t]["signed_quantity"]
                qty = float(r["quantity"])
                if r["action"] == "BUY":
                    q += qty
                else:
                    q -= qty

                state[t]["signed_quantity"] = q
                state[t]["side"] = "Short" if q < 0 else "Long"

            mode = (
                f"Latest valuation {latest['date']:%Y-%m-%d} + "
                f"submitted order {pending['date'].max():%Y-%m-%d} (estimated)"
            )

    return state, cash, mode
